import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import app


def valid_payload():
    return {
        "participant_id": "P001",
        "session_id": "12345678-abcd-4000-9000-123456789abc",
        "started_at": "2026-10-16T14:00:00.000Z",
        "finished_at": "2026-10-16T14:10:00.000Z",
        "baseline": {
            "rating_started_at": "2026-10-16T14:00:05.000Z",
            "ratings": {"valence": 4, "arousal": 5},
            "response_ms": {"valence": 1100, "arousal": 1300},
        },
        "records": [
            {
                "sequence": index,
                "rating_started_at": f"2026-10-16T14:0{index}:00.000Z",
                "ratings": {"valence": index, "arousal": 5},
                "response_ms": {"valence": 900, "arousal": 1200},
            }
            for index in range(1, 5)
        ],
    }


class ValidationTests(unittest.TestCase):
    def test_public_config_uses_two_attached_scales_and_requested_noise_segment(self):
        config = app.public_config()
        self.assertEqual(config["noise_start_seconds"], 60)
        self.assertEqual(config["noise_seconds"], 20)
        self.assertEqual(
            [item["image"] for item in config["dimensions"]],
            ["/assets/Slide8.JPG", "/assets/Slide9.JPG"],
        )

    def test_accepts_complete_experiment(self):
        cleaned = app.validate_payload(valid_payload(), complete=True)
        self.assertEqual(cleaned["participant_id"], "P001")
        self.assertEqual(cleaned["baseline"]["ratings"]["valence"], 4)
        self.assertEqual(len(cleaned["records"]), 4)
        self.assertEqual(cleaned["records"][0]["ratings"]["arousal"], 5)

    def test_rejects_formula_like_participant_id(self):
        payload = valid_payload()
        payload["participant_id"] = "=HYPERLINK('bad')"
        with self.assertRaisesRegex(ValueError, "participant_id"):
            app.validate_payload(payload, complete=True)

    def test_rejects_out_of_range_rating(self):
        payload = valid_payload()
        payload["records"][0]["ratings"]["arousal"] = 10
        with self.assertRaisesRegex(ValueError, "arousal rating"):
            app.validate_payload(payload, complete=True)

    def test_requires_every_video_on_completion(self):
        payload = valid_payload()
        payload["records"].pop()
        with self.assertRaisesRegex(ValueError, "must contain 4 records"):
            app.validate_payload(payload, complete=True)

    def test_requires_baseline_on_completion(self):
        payload = valid_payload()
        payload["baseline"] = None
        with self.assertRaisesRegex(ValueError, "pre-video baseline"):
            app.validate_payload(payload, complete=True)

    def test_requires_every_baseline_scale(self):
        payload = valid_payload()
        payload["baseline"]["ratings"]["arousal"] = None
        with self.assertRaisesRegex(ValueError, "baseline requires a response"):
            app.validate_payload(payload, complete=True)

    def test_requires_every_scale_on_completion(self):
        payload = valid_payload()
        payload["records"][0]["ratings"]["arousal"] = None
        payload["records"][0]["response_ms"]["arousal"] = None
        with self.assertRaisesRegex(ValueError, "response for every scale"):
            app.validate_payload(payload, complete=True)

    def test_allows_response_longer_than_old_timeout(self):
        payload = valid_payload()
        payload["records"][0]["response_ms"]["arousal"] = 30_000
        cleaned = app.validate_payload(payload, complete=True)
        self.assertEqual(cleaned["records"][0]["response_ms"]["arousal"], 30_000)


class WorkbookTests(unittest.TestCase):
    def test_participant_workbook_contains_ratings_and_metadata(self):
        data = app.validate_payload(valid_payload(), complete=True)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "participant.xlsx"
            app.build_participant_workbook(data, path)
            workbook = load_workbook(path, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["Response graph", "Ratings", "Session information"],
            )
            self.assertEqual(workbook.active.title, "Response graph")
            graph = workbook["Response graph"]
            self.assertEqual(graph["A2"].value, "Baseline")
            self.assertEqual(graph["B2"].value, 4)
            self.assertEqual(graph["A6"].value, "Video 4")
            self.assertEqual(graph.max_column, 3)
            self.assertEqual(len(graph._charts), 1)
            self.assertEqual(len(graph._charts[0].series), 2)
            self.assertEqual(graph._charts[0].y_axis.scaling.min, 1)
            self.assertEqual(graph._charts[0].y_axis.scaling.max, 9)
            ratings = workbook["Ratings"]
            self.assertEqual(ratings.max_row, 6)
            self.assertEqual(ratings["A2"].value, "P001")
            self.assertEqual(ratings["E2"].value, "Baseline")
            self.assertEqual(ratings["I2"].value, 4)
            self.assertEqual(ratings["K2"].value, 5)
            self.assertTrue(ratings["N2"].value)
            self.assertEqual(ratings.max_column, 14)
            self.assertEqual(ratings["E3"].value, "Post-video")
            self.assertEqual(ratings["F3"].value, 1)
            self.assertEqual(ratings["I3"].value, 1)
            self.assertEqual(ratings.freeze_panes, "A2")

    def test_master_workbook_deduplicates_session_retries(self):
        data = app.validate_payload(valid_payload(), complete=True)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "all_participants.xlsx"
            app.append_master_workbook(data, path)
            app.append_master_workbook(data, path)
            workbook = load_workbook(path)
            self.assertEqual(workbook["Ratings"].max_row, 6)

    def test_master_workbook_migrates_legacy_rows_to_post_video(self):
        data = app.validate_payload(valid_payload(), complete=True)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "all_participants.xlsx"
            from openpyxl import Workbook

            legacy = Workbook()
            sheet = legacy.active
            sheet.title = "Ratings"
            sheet.append(app.LEGACY_HEADERS)
            sheet.append(["OLD", "old-session", "start", "finish", 1] + [None] * 10)
            legacy.save(path)

            app.append_master_workbook(data, path)
            workbook = load_workbook(path)
            ratings = workbook["Ratings"]
            self.assertEqual(ratings["E2"].value, "Post-video")
            self.assertEqual(ratings["E3"].value, "Baseline")
            self.assertEqual(ratings.max_row, 7)
            self.assertIn("Legacy ratings (3 scales)", workbook.sheetnames)

    def test_master_workbook_preserves_current_three_scale_rows_during_migration(self):
        data = app.validate_payload(valid_payload(), complete=True)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "all_participants.xlsx"
            from openpyxl import Workbook

            legacy = Workbook()
            sheet = legacy.active
            sheet.title = "Ratings"
            sheet.append(app.THREE_SCALE_HEADERS)
            sheet.append(
                [
                    "OLD3",
                    "old3-session",
                    "start",
                    "finish",
                    "Baseline",
                    None,
                    "Pre-video baseline",
                    None,
                    2,
                    100,
                    3,
                    200,
                    4,
                    300,
                    "rating-start",
                    True,
                ]
            )
            legacy.save(path)

            app.append_master_workbook(data, path)
            workbook = load_workbook(path)
            ratings = workbook["Ratings"]
            archived = workbook["Legacy ratings (3 scales)"]
            self.assertEqual(archived["M2"].value, 4)
            self.assertEqual(ratings["M2"].value, "rating-start")
            self.assertTrue(ratings["N2"].value)
            self.assertEqual(ratings.max_column, 14)
            self.assertEqual(ratings.max_row, 7)


if __name__ == "__main__":
    unittest.main()
