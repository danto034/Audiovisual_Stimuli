"""Local video and Self-Assessment Manikin experiment application.

Run with::

    python app.py

The application intentionally binds only to localhost.  Its browser interface uses
the YouTube IFrame Player API and its Python backend writes participant results to
Excel workbooks in ``results/``.
"""

from __future__ import annotations

import argparse
import json
import re
import threading
import uuid
import webbrowser
from datetime import datetime, timezone
from functools import partial
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


APP_ROOT = Path(__file__).resolve().parent
WEB_ROOT = APP_ROOT / "web"
RESULTS_ROOT = APP_ROOT / "results"
CHECKPOINT_ROOT = RESULTS_ROOT / "checkpoints"

HOST = "127.0.0.1"
DEFAULT_PORT = 8787
MAX_BODY_BYTES = 1_000_000
APP_VERSION = "2026.10.17.3"

NOISE_SECONDS = 20
NOISE_START_SECONDS = 60

VIDEOS = [
    {
        "title": "The Closet | A One Minute Horror Short Film",
        "youtube_id": "JPPLHSncb-g",
        "url": "https://www.youtube.com/watch?v=JPPLHSncb-g",
    },
    {
        "title": "Sad Cinematic Emotional Intro background music 30 second [No Copyright]",
        "youtube_id": "BnKb4HzaAiM",
        "url": "https://www.youtube.com/watch?v=BnKb4HzaAiM",
        "end_seconds": 26,
    },
    {
        "title": '"Dreams Without Goals Are Just Dreams" | Denzel Washington Motivational Speech',
        "youtube_id": "yt-aM2SKLOg",
        "url": "https://www.youtube.com/watch?v=yt-aM2SKLOg",
    },
    {
        "title": "1-Minute Nature Background Sound",
        "youtube_id": "DeHUFsrCYr0",
        "url": "https://www.youtube.com/watch?v=DeHUFsrCYr0",
    },
]

NOISE_VIDEO = {
    "title": "White Noise with Black Screen",
    "youtube_id": "E82liRPcXMw",
    "url": "https://www.youtube.com/watch?v=E82liRPcXMw",
}

DIMENSIONS = (
    (
        "valence",
        "Pleasure",
        "Negative / unpleasant",
        "Positive / pleasant",
        "/assets/Slide8.JPG",
        "Pleasure scale from 1, negative and unpleasant, to 9, positive and pleasant.",
    ),
    (
        "arousal",
        "Arousal",
        "Calm / relaxed",
        "Excited / active",
        "/assets/Slide9.JPG",
        "Arousal scale from 1, calm and relaxed, to 9, excited and active.",
    ),
)

PARTICIPANT_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,63}$")
SESSION_RE = re.compile(r"^[A-Za-z0-9-]{8,64}$")
RESULT_FILE_RE = re.compile(r"^[A-Za-z0-9_.-]+\.xlsx$")
WORKBOOK_LOCK = threading.Lock()


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def public_config() -> dict[str, Any]:
    return {
        "app_version": APP_VERSION,
        "videos": VIDEOS,
        "noise_video": NOISE_VIDEO,
        "noise_seconds": NOISE_SECONDS,
        "noise_start_seconds": NOISE_START_SECONDS,
        "dimensions": [
            {
                "key": key,
                "label": label,
                "low": low,
                "high": high,
                "image": image,
                "image_alt": image_alt,
            }
            for key, label, low, high, image, image_alt in DIMENSIONS
        ],
    }


def require_string(value: Any, name: str, maximum: int = 128) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > maximum:
        raise ValueError(f"{name} is invalid")
    return value.strip()


def optional_iso_string(value: Any, name: str) -> str:
    if value is None:
        return ""
    text = require_string(value, name, 64)
    try:
        datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError(f"{name} must be an ISO date/time") from exc
    return text


def optional_rating(value: Any, name: str) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, int) or not 1 <= value <= 9:
        raise ValueError(f"{name} must be a whole number from 1 to 9")
    return value


def optional_milliseconds(value: Any, name: str) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{name} must be a number")
    rounded = int(round(value))
    if not 0 <= rounded <= 86_400_000:
        raise ValueError(f"{name} is outside the supported range")
    return rounded


def validate_payload(payload: Any, *, complete: bool) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise ValueError("Request body must be a JSON object")

    participant_id = require_string(payload.get("participant_id"), "participant_id", 64)
    if not PARTICIPANT_RE.fullmatch(participant_id):
        raise ValueError(
            "participant_id must start with a letter or number and use only letters, "
            "numbers, dots, hyphens, or underscores"
        )

    session_id = require_string(payload.get("session_id"), "session_id", 64)
    if not SESSION_RE.fullmatch(session_id):
        raise ValueError("session_id is invalid")

    started_at = optional_iso_string(payload.get("started_at"), "started_at")
    finished_at = optional_iso_string(payload.get("finished_at"), "finished_at")
    if complete and (not started_at or not finished_at):
        raise ValueError("Completed experiments require start and finish timestamps")

    raw_baseline = payload.get("baseline")
    clean_baseline: dict[str, Any] | None = None
    if raw_baseline is not None:
        if not isinstance(raw_baseline, dict):
            raise ValueError("baseline must be an object")
        baseline_ratings = raw_baseline.get("ratings")
        baseline_response_ms = raw_baseline.get("response_ms")
        if not isinstance(baseline_ratings, dict) or not isinstance(baseline_response_ms, dict):
            raise ValueError("baseline ratings and response_ms must be objects")
        clean_baseline = {
            "rating_started_at": optional_iso_string(
                raw_baseline.get("rating_started_at"), "baseline rating_started_at"
            ),
            "ratings": {
                key: optional_rating(baseline_ratings.get(key), f"baseline {key} rating")
                for key, *_ in DIMENSIONS
            },
            "response_ms": {
                key: optional_milliseconds(
                    baseline_response_ms.get(key), f"baseline {key} response_ms"
                )
                for key, *_ in DIMENSIONS
            },
        }
        if (
            any(value is None for value in clean_baseline["ratings"].values())
            or any(value is None for value in clean_baseline["response_ms"].values())
        ):
            raise ValueError("The baseline requires a response for every scale")
    if complete and clean_baseline is None:
        raise ValueError("Completed experiments require a pre-video baseline")

    records = payload.get("records")
    if not isinstance(records, list):
        raise ValueError("records must be a list")
    if complete and len(records) != len(VIDEOS):
        raise ValueError(f"A completed experiment must contain {len(VIDEOS)} records")
    if len(records) > len(VIDEOS):
        raise ValueError("Too many records")

    clean_records: list[dict[str, Any]] = []
    seen: set[int] = set()
    for raw_record in records:
        if not isinstance(raw_record, dict):
            raise ValueError("Each record must be a JSON object")
        sequence = raw_record.get("sequence")
        if isinstance(sequence, bool) or not isinstance(sequence, int):
            raise ValueError("record sequence must be a whole number")
        if not 1 <= sequence <= len(VIDEOS) or sequence in seen:
            raise ValueError("record sequence is missing, duplicated, or out of range")
        seen.add(sequence)

        ratings = raw_record.get("ratings")
        response_ms = raw_record.get("response_ms")
        if not isinstance(ratings, dict) or not isinstance(response_ms, dict):
            raise ValueError("ratings and response_ms must be objects")

        clean_record = {
            "sequence": sequence,
            "rating_started_at": optional_iso_string(
                raw_record.get("rating_started_at"), "rating_started_at"
            ),
            "ratings": {
                key: optional_rating(ratings.get(key), f"{key} rating")
                for key, *_ in DIMENSIONS
            },
            "response_ms": {
                key: optional_milliseconds(response_ms.get(key), f"{key} response_ms")
                for key, *_ in DIMENSIONS
            },
        }
        if complete and (
            any(value is None for value in clean_record["ratings"].values())
            or any(value is None for value in clean_record["response_ms"].values())
        ):
            raise ValueError("Completed experiments require a response for every scale")
        clean_records.append(clean_record)

    clean_records.sort(key=lambda item: item["sequence"])
    if complete and [item["sequence"] for item in clean_records] != list(
        range(1, len(VIDEOS) + 1)
    ):
        raise ValueError("Completed records must include every video in sequence")

    return {
        "participant_id": participant_id,
        "session_id": session_id,
        "started_at": started_at,
        "finished_at": finished_at,
        "baseline": clean_baseline,
        "records": clean_records,
    }


HEADERS = [
    "Participant ID",
    "Session ID",
    "Experiment started (UTC)",
    "Experiment finished (UTC)",
    "Measurement phase",
    "Video order",
    "Video title",
    "YouTube URL",
    "Pleasure rating (1-9)",
    "Pleasure response time (ms)",
    "Arousal rating (1-9)",
    "Arousal response time (ms)",
    "Rating period started (UTC)",
    "Response required before advancement",
]

THREE_SCALE_HEADERS = [
    "Participant ID",
    "Session ID",
    "Experiment started (UTC)",
    "Experiment finished (UTC)",
    "Measurement phase",
    "Video order",
    "Video title",
    "YouTube URL",
    "Pleasure rating (1-9)",
    "Pleasure response time (ms)",
    "Arousal rating (1-9)",
    "Arousal response time (ms)",
    "Dominance rating (1-9)",
    "Dominance response time (ms)",
    "Rating period started (UTC)",
    "Response required before advancement",
]

LEGACY_HEADERS = [
    "Participant ID",
    "Session ID",
    "Experiment started (UTC)",
    "Experiment finished (UTC)",
    "Video order",
    "Video title",
    "YouTube URL",
    "Affect / valence rating (1-9)",
    "Affect / valence response time (ms)",
    "Arousal rating (1-9)",
    "Arousal response time (ms)",
    "Dominance rating (1-9)",
    "Dominance response time (ms)",
    "Rating period started (UTC)",
    "Response required before advancement",
]


def result_rows(data: dict[str, Any]) -> list[list[Any]]:
    baseline = data["baseline"]
    rows: list[list[Any]] = [
        [
            data["participant_id"],
            data["session_id"],
            data["started_at"],
            data["finished_at"],
            "Baseline",
            None,
            "Pre-video baseline",
            None,
            baseline["ratings"]["valence"],
            baseline["response_ms"]["valence"],
            baseline["ratings"]["arousal"],
            baseline["response_ms"]["arousal"],
            baseline["rating_started_at"],
            True,
        ]
    ]
    for record in data["records"]:
        video = VIDEOS[record["sequence"] - 1]
        rows.append(
            [
                data["participant_id"],
                data["session_id"],
                data["started_at"],
                data["finished_at"],
                "Post-video",
                record["sequence"],
                video["title"],
                video["url"],
                record["ratings"]["valence"],
                record["response_ms"]["valence"],
                record["ratings"]["arousal"],
                record["response_ms"]["arousal"],
                record["rating_started_at"],
                True,
            ]
        )
    return rows


def style_ratings_sheet(sheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="19324A")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 42

    widths = [16, 38, 24, 24, 18, 12, 54, 48, 22, 27, 20, 25, 25, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_response_graph_sheet(workbook: Workbook, data: dict[str, Any]) -> None:
    """Add a combined, editable Excel line chart for the five measurement timepoints."""
    sheet = workbook.create_sheet("Response graph", 0)
    sheet.sheet_view.showGridLines = False
    sheet.append(["Timepoint", "Pleasure", "Arousal"])

    baseline = data["baseline"]
    sheet.append(
        [
            "Baseline",
            baseline["ratings"]["valence"],
            baseline["ratings"]["arousal"],
        ]
    )
    for record in data["records"]:
        sheet.append(
            [
                f"Video {record['sequence']}",
                record["ratings"]["valence"],
                record["ratings"]["arousal"],
            ]
        )

    header_fill = PatternFill("solid", fgColor="19324A")
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2, max_row=6, min_col=1, max_col=3):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(
                horizontal="left" if cell.column == 1 else "right", vertical="center"
            )
    sheet.column_dimensions["A"].width = 18
    for column in ("B", "C"):
        sheet.column_dimensions[column].width = 14
    sheet["A8"] = "Scale range: 1 to 9"
    sheet["A8"].font = Font(name="Arial", size=10, italic=True, color="53636D")

    chart = LineChart()
    chart.title = "SAM responses over time"
    chart.style = 13
    chart.height = 10
    chart.width = 18
    chart.legend.position = "t"
    chart.x_axis.title = "Timepoint"
    chart.y_axis.title = "Rating (1–9)"
    chart.y_axis.scaling.min = 1
    chart.y_axis.scaling.max = 9
    chart.y_axis.majorUnit = 1
    chart.add_data(Reference(sheet, min_col=2, max_col=3, min_row=1, max_row=6), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=6))

    colors = ("2A9D8F", "E76F51")
    for series, color in zip(chart.series, colors):
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 28575
        series.marker.symbol = "circle"
        series.marker.size = 7
        series.marker.graphicalProperties.solidFill = color
        series.marker.graphicalProperties.line.solidFill = color

    sheet.add_chart(chart, "F2")


def build_participant_workbook(data: dict[str, Any], path: Path) -> None:
    workbook = Workbook()
    ratings = workbook.active
    ratings.title = "Ratings"
    ratings.append(HEADERS)
    for row in result_rows(data):
        ratings.append(row)
    style_ratings_sheet(ratings)

    build_response_graph_sheet(workbook, data)
    workbook.active = 0

    session = workbook.create_sheet("Session information")
    session.append(["Field", "Value"])
    details = [
        ("Participant ID", data["participant_id"]),
        ("Session ID", data["session_id"]),
        ("Experiment started (UTC)", data["started_at"]),
        ("Experiment finished (UTC)", data["finished_at"]),
        ("Export created (UTC)", utc_now_iso()),
        ("Video sequence", "Fixed; 1 through 4"),
        ("Video 2 playback window", "0:00 through 0:26; remaining content omitted"),
        ("Pre-video baseline", "Pleasure and arousal; responses are recorded"),
        ("Scale advancement rule", "A response is required; there is no auto-advance"),
        ("White-noise interval after videos 1-3 (seconds)", NOISE_SECONDS),
        ("White-noise source starts at (seconds)", NOISE_START_SECONDS),
        ("1 - Pleasure", "Negative / unpleasant"),
        ("9 - Pleasure", "Positive / pleasant"),
        ("1 - Arousal", "Calm / relaxed"),
        ("9 - Arousal", "Excited / active"),
    ]
    for key, value in details:
        session.append([key, value])
    session.column_dimensions["A"].width = 48
    session.column_dimensions["B"].width = 64
    for cell in session[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="19324A")
    for row in session.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook.save(path)


def append_master_workbook(data: dict[str, Any], path: Path) -> None:
    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook["Ratings"]
        existing_headers = [
            sheet.cell(row=1, column=column).value
            for column in range(1, sheet.max_column + 1)
        ]
        if existing_headers in (LEGACY_HEADERS, THREE_SCALE_HEADERS):
            archived_title = "Legacy ratings (3 scales)"
            suffix = 2
            while archived_title in workbook.sheetnames:
                archived_title = f"Legacy ratings (3 scales) {suffix}"
                suffix += 1
            archived = workbook.copy_worksheet(sheet)
            archived.title = archived_title

        if existing_headers == LEGACY_HEADERS:
            sheet.insert_cols(5)
            for row_number in range(2, sheet.max_row + 1):
                sheet.cell(row=row_number, column=5, value="Post-video")
            sheet.delete_cols(13, 2)
            for column, heading in enumerate(HEADERS, start=1):
                sheet.cell(row=1, column=column, value=heading)
        elif existing_headers == THREE_SCALE_HEADERS:
            sheet.delete_cols(13, 2)
            for column, heading in enumerate(HEADERS, start=1):
                sheet.cell(row=1, column=column, value=heading)
        elif existing_headers != HEADERS:
            raise ValueError(
                "all_participants.xlsx has an incompatible column layout; "
                "preserve it under another name before starting a new session"
            )
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Ratings"
        sheet.append(HEADERS)

    # A browser can retry a save after a lost HTTP response.  Treat the session ID
    # as an idempotency key so a retry never duplicates rows in the master workbook.
    existing_sessions = {
        cell.value for cell in sheet["B"][1:] if isinstance(cell.value, str)
    }
    if data["session_id"] not in existing_sessions:
        for row in result_rows(data):
            sheet.append(row)
    style_ratings_sheet(sheet)

    temporary_path = path.with_name(f".{path.stem}-{uuid.uuid4().hex}.tmp.xlsx")
    try:
        workbook.save(temporary_path)
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)


def save_checkpoint(data: dict[str, Any], *, status: str) -> Path:
    CHECKPOINT_ROOT.mkdir(parents=True, exist_ok=True)
    checkpoint = CHECKPOINT_ROOT / f"{data['session_id']}.json"
    document = {**data, "status": status, "server_updated_at": utc_now_iso()}
    temporary = checkpoint.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(document, indent=2), encoding="utf-8")
    temporary.replace(checkpoint)
    return checkpoint


def finish_experiment(data: dict[str, Any]) -> tuple[Path, str | None]:
    RESULTS_ROOT.mkdir(parents=True, exist_ok=True)
    started = datetime.fromisoformat(data["started_at"].replace("Z", "+00:00"))
    timestamp = started.strftime("%Y%m%d_%H%M%S")
    filename = f"{data['participant_id']}_{timestamp}_{data['session_id'][:8]}.xlsx"
    participant_path = RESULTS_ROOT / filename
    master_warning: str | None = None

    with WORKBOOK_LOCK:
        build_participant_workbook(data, participant_path)
        try:
            append_master_workbook(data, RESULTS_ROOT / "all_participants.xlsx")
        except PermissionError:
            master_warning = (
                "The participant workbook was saved, but all_participants.xlsx could not be "
                "updated because it appears to be open in Excel."
            )
        except ValueError as exc:
            master_warning = (
                "The participant workbook was saved, but the cumulative workbook was not "
                f"updated: {exc}"
            )
        save_checkpoint(data, status="complete")
    return participant_path, master_warning


class ExperimentHandler(SimpleHTTPRequestHandler):
    server_version = "SAMExperiment/1.0"

    def end_headers(self) -> None:
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "strict-origin-when-cross-origin")
        self.send_header("Cache-Control", "no-store")
        super().end_headers()

    def log_message(self, fmt: str, *args: Any) -> None:
        print(f"[{self.log_date_time_string()}] {fmt % args}")

    def send_json(self, status: int, document: dict[str, Any]) -> None:
        body = json.dumps(document).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def read_json(self) -> Any:
        content_type = self.headers.get("Content-Type", "")
        if "application/json" not in content_type:
            raise ValueError("Content-Type must be application/json")
        try:
            size = int(self.headers.get("Content-Length", "0"))
        except ValueError as exc:
            raise ValueError("Content-Length is invalid") from exc
        if not 0 < size <= MAX_BODY_BYTES:
            raise ValueError("Request body is empty or too large")
        try:
            return json.loads(self.rfile.read(size).decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError("Request body is not valid JSON") from exc

    def do_GET(self) -> None:  # noqa: N802 - inherited HTTP method name
        route = unquote(urlparse(self.path).path)
        if route == "/api/config":
            self.send_json(HTTPStatus.OK, public_config())
            return
        if route.startswith("/api/download/"):
            filename = route.removeprefix("/api/download/")
            if not RESULT_FILE_RE.fullmatch(filename):
                self.send_error(HTTPStatus.NOT_FOUND)
                return
            path = RESULTS_ROOT / filename
            if not path.is_file():
                self.send_error(HTTPStatus.NOT_FOUND)
                return
            content = path.read_bytes()
            self.send_response(HTTPStatus.OK)
            self.send_header(
                "Content-Type",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
            self.send_header("Content-Length", str(len(content)))
            self.end_headers()
            self.wfile.write(content)
            return
        if route == "/":
            self.path = "/index.html"
        super().do_GET()

    def do_POST(self) -> None:  # noqa: N802 - inherited HTTP method name
        route = urlparse(self.path).path
        try:
            if route == "/api/checkpoint":
                data = validate_payload(self.read_json(), complete=False)
                save_checkpoint(data, status="in_progress")
                self.send_json(HTTPStatus.OK, {"ok": True})
                return
            if route == "/api/finish":
                data = validate_payload(self.read_json(), complete=True)
                path, warning = finish_experiment(data)
                self.send_json(
                    HTTPStatus.OK,
                    {
                        "ok": True,
                        "filename": path.name,
                        "download_url": f"/api/download/{path.name}",
                        "saved_path": str(path),
                        "master_path": str(RESULTS_ROOT / "all_participants.xlsx"),
                        "warning": warning,
                    },
                )
                return
            self.send_json(HTTPStatus.NOT_FOUND, {"ok": False, "error": "Not found"})
        except ValueError as exc:
            print(f"Request rejected on {route}: {exc}")
            self.send_json(HTTPStatus.BAD_REQUEST, {"ok": False, "error": str(exc)})
        except Exception as exc:  # Keep the experiment page alive and report a useful error.
            print(f"Save error: {exc!r}")
            self.send_json(
                HTTPStatus.INTERNAL_SERVER_ERROR,
                {"ok": False, "error": "The results could not be saved. Please try again."},
            )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run the video and SAM rating experiment")
    parser.add_argument("--port", type=int, default=DEFAULT_PORT, help="Local port to use")
    parser.add_argument(
        "--no-browser", action="store_true", help="Do not open the browser automatically"
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if not WEB_ROOT.is_dir():
        raise SystemExit(f"Web interface not found: {WEB_ROOT}")
    if not 1 <= args.port <= 65535:
        raise SystemExit("Port must be between 1 and 65535")

    handler = partial(ExperimentHandler, directory=str(WEB_ROOT))
    try:
        server = ThreadingHTTPServer((HOST, args.port), handler)
    except OSError as exc:
        raise SystemExit(f"Could not start on http://{HOST}:{args.port}: {exc}") from exc

    url = f"http://{HOST}:{args.port}"
    print("SAM video experiment is ready.")
    print(f"Open: {url}")
    print(f"Results: {RESULTS_ROOT}")
    print("Press Ctrl+C in this window to stop the application.")
    if not args.no_browser:
        threading.Timer(0.8, lambda: webbrowser.open(url)).start()

    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopping application.")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
